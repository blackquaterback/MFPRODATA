import os
import pandas as pd
from openpyxl import load_workbook

class CanaraMutualFundParser:
    def __init__(self, input_dir, amc_name="Canara Robeco Mutual Fund"):
        self.input_dir = input_dir
        self.amc_name = amc_name

    def clean_scheme_name(self, raw_text):
        if not isinstance(raw_text, str):
            return "Unknown Scheme"
        return raw_text.split("(")[0].split("-")[0].strip()

    def parse_sheet(self, file_path, sheet_name, scheme_name):
        try:
            df_raw = pd.read_excel(file_path, sheet_name=sheet_name, header=3)
            df_raw.dropna(how="all", inplace=True)
            df_raw.columns = df_raw.columns.str.strip().str.replace(r"\s+", " ", regex=True)

            df_raw.rename(columns={
                "Name of the Instrument": "Name of Instrument",
                "Market/Fair Value (Rs. in Lacs)": "Market Value",
                "% to Net Assets": "% to Net Assets",
                "Yield %": "Yield",
                "Industry / Rating": "Industry",
                "Market Capitalization": "Market Capitalization"
            }, inplace=True)

            if "ISIN" in df_raw.columns:
                isin_idx = df_raw.columns.get_loc("ISIN")
                df_raw.insert(isin_idx + 1, "Coupon", "")
            else:
                df_raw["Coupon"] = ""

            if "% to Net Assets" in df_raw.columns:
                df_raw["% to Net Assets"] = (
                    df_raw["% to Net Assets"]
                    .astype(str)
                    .str.replace(r"[^\d.\-]", "", regex=True)
                    .replace("", "0")
                    .astype(float)
                )

            if "Yield" in df_raw.columns:
                df_raw["Yield"] = (
                    df_raw["Yield"]
                    .astype(str)
                    .str.replace(r"[^\d.\-]", "", regex=True)
                    .replace("", "0")
                    .astype(float)
                )
            else:
                df_raw["Yield"] = 0.0

            df_raw["Yield"] = df_raw["Yield"].replace(0, "")
            df_raw["Type"] = df_raw["Yield"].apply(lambda x: "Debt" if pd.notna(x) and x != "" else "Equity")
            df_raw["Scheme Name"] = scheme_name
            df_raw["AMC"] = self.amc_name

            required_columns = [
                "Name of Instrument", "ISIN", "Coupon", "Industry", "Quantity",
                "Market Value", "% to Net Assets", "Market Capitalization", "Yield",
                "Type", "Scheme Name", "AMC"
            ]

            for col in required_columns:
                if col not in df_raw.columns:
                    df_raw[col] = ""

            df_final = df_raw[required_columns]

            if "Market Capitalization" in df_final.columns:
                df_final.drop(columns=["Market Capitalization"], inplace=True)

            df_final = df_final[df_final["ISIN"].astype(str).str.startswith("IN", na=False)]

            return df_final

        except Exception as e:
            print(f"Skipping sheet '{sheet_name}' in '{os.path.basename(file_path)}' due to error: {e}")
            return None

    def run(self):
        all_data = []

        for filename in os.listdir(self.input_dir):
            if filename.endswith(".xlsx"):
                file_path = os.path.join(self.input_dir, filename)
                try:
                    wb = load_workbook(file_path, data_only=True)
                    for sheet in wb.sheetnames:
                        ws = wb[sheet]
                        raw_scheme = str(ws["B1"].value).strip()
                        scheme_name = self.clean_scheme_name(raw_scheme)

                        df = self.parse_sheet(file_path, sheet, scheme_name)
                        if df is not None:
                            all_data.append(df)

                except Exception as e:
                    print(f"Skipping file '{filename}' due to error: {e}")

        if all_data:
            return pd.concat(all_data, ignore_index=True)
        return pd.DataFrame()

def save_to_excel(df, output_path):
    df.to_excel(output_path, index=False)
    print(f"Cleaned and combined file saved at: {output_path}")

input_folder = "C:/Users/SAURABH/Desktop/Portfolio/data/Canara Robeco Mutual Fund"
output_file = "canara_robeco_mutualfund.xlsx"
parser = CanaraMutualFundParser(input_folder)
result_df = parser.run()
if not result_df.empty:
    save_to_excel(result_df, output_file)
else:
    print("No valid data extracted.")
