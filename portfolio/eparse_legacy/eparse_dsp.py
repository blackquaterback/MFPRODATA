import os
import pandas as pd
from openpyxl import load_workbook

file_path = "C:\Users\SAURABH\Desktop\Portfolio\data\DSP Mutual Fund\dsp-isin-debt-portfolio-as-on-28-feb-2025.xlsx"
output_path = "dsp_mutualfund_portfolio.xlsx"

class DSPMutualFundParser:
    def __init__(self, file_path):
        self.file_path = file_path
        self.amc_name = "DSP"

    def clean_scheme_name(self, name):
        if not isinstance(name, str):
            return "Unknown Scheme"
        return name.split("(")[0].split("-")[0].strip()

    def process_sheet(self, sheet, wb):
        try:
            ws = wb[sheet]
            raw_name = ws["B1"].value
            scheme_name = self.clean_scheme_name(str(raw_name).strip())

            df = pd.read_excel(self.file_path, sheet_name=sheet, header=3)
            df.dropna(how="all", inplace=True)

            if "Rating/Industry" in df.columns:
                df.rename(columns={"Rating/Industry": "Industry"}, inplace=True)
            elif "Industry / Rating" in df.columns:
                df.rename(columns={"Industry / Rating": "Industry"}, inplace=True)
            elif "Rating" in df.columns:
                df.rename(columns={"Rating": "Industry"}, inplace=True)

            df.rename(columns={
                "Name of the Instrument": "Name of Instrument",
                "Market value (Rs. In lakhs)": "Market Value",
                "Market/Fair Value(Rs. in Lacs)": "Market Value",
                "% to Net Assets": "% to Net Assets",
                "Rounded % to Net Assets": "% to Net Assets",
                "YTM (%)": "Yield",
                "YTM~": "Yield",
                "YTM": "Yield"
            }, inplace=True)

            if "ISIN" in df.columns:
                idx = df.columns.get_loc("ISIN")
                df.insert(idx + 1, "Coupon", "")
            else:
                df["Coupon"] = ""

            if "% to Net Assets" in df.columns:
                df["% to Net Assets"] = (
                    pd.to_numeric(
                        df["% to Net Assets"].astype(str).str.replace(r"[^\d.\-]", "", regex=True),
                        errors="coerce"
                    ) * 100
                )

            if "Yield" in df.columns:
                df["Yield"] = (
                    pd.to_numeric(
                        df["Yield"].astype(str).str.replace(r"[^\d.\-]", "", regex=True),
                        errors="coerce"
                    ) * 100
                )

            df["Type"] = df["Yield"].apply(lambda x: "Debt" if pd.notna(x) and x != 0 else "Equity")
            df["Scheme Name"] = scheme_name
            df["AMC"] = self.amc_name

            required_columns = [
                "Name of Instrument", "ISIN", "Coupon", "Industry", "Quantity",
                "Market Value", "% to Net Assets", "Yield", "Type", "Scheme Name", "AMC"
            ]

            for col in required_columns:
                if col not in df.columns:
                    df[col] = ""

            df = df[
                df["ISIN"].astype(str).str.startswith("IN") |
                df["Name of Instrument"].astype(str).str.contains("Clearing Corporation of India Ltd", na=False)
            ]

            return df[required_columns]

        except Exception as e:
            print(f"Skipping sheet '{sheet}' due to error: {e}")
            return None

    def run(self):
        wb = load_workbook(self.file_path, data_only=True)
        all_data = []
        for sheet in wb.sheetnames:
            df = self.process_sheet(sheet, wb)
            if df is not None:
                all_data.append(df)
        if all_data:
            return pd.concat(all_data, ignore_index=True)
        return pd.DataFrame()

def save_to_excel(df, output_path):
    df.to_excel(output_path, index=False)
    print(f"Cleaned file saved at: {output_path}")

file_path = "./dsp-isin-debt-portfolio-as-on-28-feb-2025.xlsx"
output_path = "./dsp_mutualfund.xlsx"
parser = DSPMutualFundParser(file_path)
result = parser.run()
  if not result.empty:
    save_to_excel(result, output_path)
  else:
    print("No valid data extracted.")
